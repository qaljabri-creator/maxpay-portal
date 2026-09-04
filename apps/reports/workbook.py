"""Turning a report into an ``.xlsx`` file.

Two sheets, always in the same order:

1. **الملخّص** — what was asked for, and what it came to. An exported file
   outlives the screen it was taken from, so it states its own question:
   a spreadsheet of numbers with no record of what was filtered is one
   somebody reads the wrong way three months later.
2. **الطلبات** — one row per request.

**The guard.** :func:`build` refuses to write a single cell of a merchant's
workbook until the rows have been through
:func:`apps.merchant_panel.anonymity.assert_anonymous`. That is the same check
that guards every other merchant-facing byte, applied here because an export is
the one artefact that leaves the system entirely: a screen can be closed, a JSON
response is gone when the tab is, and a file is forwarded. It fails closed — a
leak becomes a 500, which is an incident, rather than a file, which is a breach.

It is belt and braces rather than the only protection: the merchant surface
builds its rows from the same masked serializers the screens use, so by the time
they reach here they have already passed the three layers in
:mod:`apps.merchant_panel.anonymity`. This is the fourth, and it is here because
the cost of being wrong is different for a file.

Written with ``openpyxl`` in ``write_only`` mode. A report over a year of
requests should stream rows out rather than assemble the whole sheet in memory
first, and nothing here needs to read a cell back.
"""

import io
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from django.utils import timezone
from django.utils.dateparse import parse_datetime
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

#: Black ground, one accent. The identity survives into the file, because a
#: report is a thing people put in front of other people.
HEADER_FILL = PatternFill("solid", fgColor="FF001A1A")
HEADER_FONT = Font(name="Cairo", bold=True, color="FF19FFFD")
BODY_FONT = Font(name="Cairo")
TITLE_FONT = Font(name="Cairo", bold=True, size=14)

#: Excel has no "Iraqi dinar" built in, and a raw float reads as a phone number
#: at these magnitudes.
#:
#: The dinar carries no decimals — the fils is out of circulation, and a figure
#: showing one is showing a denomination that does not exist. `.##` rather than
#: nothing, so a request settled before that rule still displays the half-dinar
#: it actually recorded instead of being silently restated.
IQD_FORMAT = "#,##0.##"
#: A rate is a ratio, not an amount, and Finance may legitimately set 1470.25.
RATE_FORMAT = "#,##0.##"
#: The cent is real money in the currency the client's balance is held in.
USD_FORMAT = "#,##0.00"


@dataclass(frozen=True)
class Column:
    """One column of the rows sheet.

    ``key`` is read out of the row dict, so the column set and the row builder
    cannot drift: a column naming a key nobody produces comes out empty and
    visibly so, rather than silently shifting every value one place left.
    """

    key: str
    label: str
    width: int = 18
    number_format: str = ""
    #: How many places the *screen* preview shows. The workbook uses
    #: ``number_format`` and lets Excel do it; a preview has no Excel, so the
    #: two are stated separately rather than one being parsed out of the other.
    decimals: int = 2


def _as_datetime(value):
    """A timestamp, whatever shape it arrived in.

    The two surfaces hand their rows over differently and both are right to.
    Finance reads model fields, so a timestamp is a ``datetime``. A merchant's
    rows come out of the masked serializers — the whole point being that
    nothing bypasses them — and a serializer produces an ISO **string**.

    Parsed here rather than converted in either caller, because the
    alternative is a column of text where a reader expects a date: unsortable
    in Excel, and three lines of ``2026-08-24T10:38:49.407390+00:00`` on the
    screen.
    """
    if isinstance(value, datetime):
        return value
    if isinstance(value, str) and len(value) >= 19 and value[4] == "-" and "T" in value:
        return parse_datetime(value)
    return None


def _as_number(value):
    """A figure, whatever shape it arrived in, or ``None`` if it is not one.

    DRF renders a ``DecimalField`` as a **string** so no precision is lost on
    the way out, which is right for JSON and wrong for a spreadsheet: a column
    of text is a column nobody can sum. Only the columns that declared a
    ``number_format`` are put through this — a wallet number is digits too, and
    coercing `07700000001` to a number would silently eat its leading zero.
    """
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, str):
        try:
            return float(Decimal(value.strip()))
        except (InvalidOperation, ValueError):
            return None
    return None


def _value(value, *, numeric: bool = False):
    """Everything openpyxl will not take, turned into something it will."""
    if value is None:
        return ""
    if numeric:
        number = _as_number(value)
        if number is not None:
            return number
    if isinstance(value, Decimal):
        return float(value)
    stamp = _as_datetime(value)
    if stamp is not None:
        # Naive and local: Excel has no timezone, and a reader in Baghdad
        # comparing a UTC stamp against their own clock is a support ticket.
        return (
            timezone.localtime(stamp).replace(tzinfo=None)
            if timezone.is_aware(stamp)
            else stamp
        )
    if isinstance(value, (date, str, int, float, bool)):
        return value
    return str(value)


def display(value, *, numeric: bool = False, decimals: int = 2) -> str:
    """A value as the screen should read it.

    The workbook keeps the raw object — Excel wants a real number to sum and a
    real date to sort. The screen cannot: a `Decimal` rendered by Django under
    an Arabic locale comes out `101850,00`, with a comma where every other
    figure in the panel has a dot, and a `datetime` comes out as three lines of
    long-form month name inside a table cell eighty pixels wide.
    """
    if value is None or value == "":
        return ""
    if numeric:
        number = _as_number(value)
        if number is not None:
            if decimals == 0:
                # Trailing zeros hidden, a real fraction kept: the same rule the
                # panels apply to a dinar figure.
                return f"{number:,.2f}".rstrip("0").rstrip(".")
            return f"{number:,.{decimals}f}"
    if isinstance(value, Decimal):
        quantised = value.quantize(Decimal("0.01"))
        return f"{quantised:,.2f}"
    stamp = _as_datetime(value)
    if stamp is not None:
        local = timezone.localtime(stamp) if timezone.is_aware(stamp) else stamp
        return local.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def cells_for(rows: list[dict], columns) -> list[list[dict]]:
    """Rows as cells already paired with their column, formatted for a screen.

    Django templates cannot look a key up by a variable, and the alternative —
    a `dictkey` filter — would put the column/row alignment in the template
    where nothing checks it. Built here, a column naming a key nobody produces
    comes out visibly empty instead of silently shifting the row.
    """
    cells = []
    for row in rows:
        line = []
        for column in columns:
            raw = row.get(column.key)
            numeric = bool(column.number_format)
            # A figure or a timestamp is a Latin string inside an RTL cell, and
            # without isolation the bidi algorithm reorders it: `2026-08-24
            # 13:48` comes out as `2026-08- 13:48 24`. The template turns this
            # into the `.num` class the panels already use for exactly that.
            line.append({
                "label": column.label,
                "value": display(raw, numeric=numeric, decimals=column.decimals),
                "key": column.key,
                "ltr": numeric or _as_datetime(raw) is not None,
            })
        cells.append(line)
    return cells


def build(
    *,
    title: str,
    filters: list[tuple[str, str]],
    summary_rows: list[tuple[str, object]],
    tables: list[tuple[str, list[str], list[list]]],
    columns: list[Column],
    rows: list[dict],
    masked: bool,
) -> bytes:
    """Render one report as an ``.xlsx`` byte string.

    ``masked`` is not a style flag. It is the caller stating which surface this
    file is leaving by, and when it is true nothing is written until the rows
    have proved they carry no client identity.
    """
    if masked:
        # Before the workbook exists, not after: there is no half-written file
        # to leak if this raises.
        from apps.merchant_panel.anonymity import assert_anonymous

        assert_anonymous(rows, where="merchant report export")
        for column in columns:
            from apps.merchant_panel.anonymity import AnonymityError, is_forbidden_key

            if is_forbidden_key(column.key):
                raise AnonymityError(
                    f"merchant report export would carry client identity in column "
                    f"{column.key!r} (spec §2)."
                )

    book = Workbook(write_only=True)

    _summary_sheet(book, title, filters, summary_rows, tables)
    _rows_sheet(book, columns, rows)

    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def _summary_sheet(book, title, filters, summary_rows, tables) -> None:
    sheet = book.create_sheet(str(_("الملخّص")))
    sheet.sheet_view.rightToLeft = True
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 22
    for letter in "CDE":
        sheet.column_dimensions[letter].width = 18

    sheet.append([_cell(sheet, title, TITLE_FONT)])
    sheet.append([])

    sheet.append([_cell(sheet, str(_("التصفية")), HEADER_FONT, HEADER_FILL)])
    for label, value in filters:
        sheet.append([_cell(sheet, label, BODY_FONT), _cell(sheet, value, BODY_FONT)])
    sheet.append([])

    sheet.append([_cell(sheet, str(_("الإجمالي")), HEADER_FONT, HEADER_FILL)])
    for label, value in summary_rows:
        sheet.append([
            _cell(sheet, label, BODY_FONT),
            _cell(sheet, _value(value), BODY_FONT),
        ])
    sheet.append([])

    for caption, headers, body in tables:
        sheet.append([_cell(sheet, caption, HEADER_FONT, HEADER_FILL)])
        sheet.append([_cell(sheet, head, HEADER_FONT, HEADER_FILL) for head in headers])
        for line in body:
            sheet.append([_cell(sheet, _value(item), BODY_FONT) for item in line])
        sheet.append([])

    sheet.append([
        _cell(sheet, str(_("وقت التصدير")), BODY_FONT),
        _cell(sheet, _value(timezone.now()), BODY_FONT),
    ])


def _rows_sheet(book, columns: list[Column], rows: list[dict]) -> None:
    sheet = book.create_sheet(str(_("الطلبات")))
    sheet.sheet_view.rightToLeft = True
    # Frozen under the header, so a thousand rows still say what each column is.
    sheet.freeze_panes = "A2"

    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = column.width

    sheet.append([_cell(sheet, column.label, HEADER_FONT, HEADER_FILL) for column in columns])

    for row in rows:
        sheet.append([
            _cell(
                sheet,
                _value(row.get(column.key), numeric=bool(column.number_format)),
                BODY_FONT,
                number_format=column.number_format,
            )
            for column in columns
        ])


def _cell(sheet, value, font, fill=None, number_format: str = ""):
    """A write-only cell carrying its own formatting.

    A ``write_only`` workbook has no addressable cells, so styling travels with
    the value instead of being applied to a range afterwards. The sheet is a
    parameter rather than module state because Django serves requests on
    threads, and two reports building at once must not share a handle.
    """
    cell = WriteOnlyCell(sheet, value=value)
    cell.font = font
    if fill is not None:
        cell.fill = fill
    if number_format:
        cell.number_format = number_format
    cell.alignment = Alignment(horizontal="right", readingOrder=2)
    return cell
